from pathlib import Path
import csv, re, json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).parent
source = BASE / '课堂试点与评价说明.md'
DOCX = BASE / '课堂试点与评价说明.docx'

def font(style, east, size, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    fonts=rpr.find(qn('w:rFonts'))
    if fonts is None: fonts=OxmlElement('w:rFonts'); rpr.insert(0,fonts)
    fonts.set(qn('w:eastAsia'),east)
    fonts.set(qn('w:ascii'),'Times New Roman')
    fonts.set(qn('w:hAnsi'),'Times New Roman')
    for a in ['w:eastAsiaTheme','w:asciiTheme','w:hAnsiTheme']: fonts.attrib.pop(qn(a),None)

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
sec.footer_distance=Cm(1.25)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal=doc.styles['Normal'];font(normal,'SimSun',12)
pf=normal.paragraph_format
pf.first_line_indent=Pt(24);pf.line_spacing=Pt(22);pf.space_after=Pt(4);pf.widow_control=True
for st in ['Title','Heading 1','Heading 2']:
    s=doc.styles[st];font(s,'SimHei',18 if st=='Title' else (14 if st=='Heading 1' else 12))
    s.paragraph_format.first_line_indent=Pt(0)
    s.paragraph_format.line_spacing=Pt(22 if st!='Title' else 28)
    s.paragraph_format.keep_with_next=True;s.paragraph_format.keep_together=True
    s.paragraph_format.space_before=Pt(10 if st!='Title' else 0)
    s.paragraph_format.space_after=Pt(6 if st!='Title' else 12)
font(doc.styles['Footer'],'SimSun',10)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='1';rr.append(tt);field.append(rr);footer._p.append(field)

def table(rows):
    coln=len(rows[0]);tb=doc.add_table(rows=1, cols=coln);tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    widths=[2.1,6.95,6.95] if coln==3 else [3.4,12.6]
    for column,width in zip(tb.columns,widths): column.width=Cm(width)
    pr=tb._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for tag,value in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
        e=OxmlElement('w:'+tag);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for i,row in enumerate(rows):
        cells=tb.rows[0].cells if i==0 else tb.add_row().cells
        trPr=cells[0]._tc.getparent().get_or_add_trPr()
        cannot=OxmlElement('w:cantSplit');trPr.append(cannot)
        if i==0: repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for j,text in enumerate(row):
            cell=cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=Pt(22);p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==0 or coln==3 else WD_ALIGN_PARAGRAPH.LEFT
            run=p.add_run(text);run.bold=i==0
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7');cell._tc.get_or_add_tcPr().append(shade)
    after=doc.add_paragraph();after.paragraph_format.space_after=Pt(0);after.paragraph_format.line_spacing=Pt(4)

lines=source.read_text().splitlines(); i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',x) for x in vals):rows.append(vals)
            i+=1
        table(rows);continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif re.match(r'^[一二三四五六七八九十]+、',line):doc.add_paragraph(line,'Heading 1')
    else:doc.add_paragraph(line)
    i+=1

doc.core_properties.title='专利情报AI工具课堂试点与评价说明'
doc.core_properties.subject='45分钟课堂方案与未实施评价模板'
doc.core_properties.author='专利情报AI项目'
doc.core_properties.keywords='课堂试点,专利情报,评价量表,待实施'
doc.save(DOCX)
fields={
 'task_records.csv':['participant_id','experience_level','sequence_group','task_code','condition','data_mode','inference_mode','dataset_version','query_version','product_version','started_at','ended_at','elapsed_seconds','interruption_seconds','net_seconds','timeout','completion_status','source_open_count','valid_evidence_count','invalid_evidence_count','operation_help_count','content_hint_count','critical_error_codes','artifact_path','notes'],
 'score_records.csv':['participant_id','task_code','condition','artifact_path','rater_id','rubric_version','efficiency_0_4','source_traceability_0_4','evidence_support_0_4','scope_awareness_0_4','evidence_total_0_12','transfer_topic_0_4','transfer_query_0_4','transfer_relevance_0_4','transfer_scope_0_4','transfer_total_0_16','score_basis','is_consensus','scored_at'],
 'feedback_records.csv':['participant_id','task_code','step','observed_behavior','participant_quote','observer_interpretation','use_again_intent','use_again_reason','alternative_tool','issue_category','artifact_reference','followup_action','recorded_at']}
for name,header in fields.items():
    if not (BASE/name).exists():
        with (BASE/name).open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerow(header)
(BASE/'qa'/'fields.json').write_text(json.dumps(fields,ensure_ascii=False,indent=2))
print(DOCX)
print('CSV templates contain headers only; no participant rows.')
