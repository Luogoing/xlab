from pathlib import Path
import csv, re, json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).parent
source = BASE / '产品与技术说明.md'
DOCX = BASE / '产品与技术说明.docx'

def font(style, east, size, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    fonts=rpr.find(qn('w:rFonts'))
    if fonts is None: fonts=OxmlElement('w:rFonts'); rpr.insert(0,fonts)
    fonts.set(qn('w:eastAsia'),east)
    fonts.set(qn('w:ascii'),'Times New Roman')
    fonts.set(qn('w:hAnsi'),'Times New Roman')
    for a in ['w:eastAsiaTheme','w:asciiTheme','w:hAnsiTheme']: fonts.attrib.pop(qn(a),None)

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
sec.footer_distance=Cm(1.25)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal=doc.styles['Normal'];font(normal,'SimSun',12)
pf=normal.paragraph_format
pf.first_line_indent=Pt(24);pf.line_spacing=Pt(22);pf.space_after=Pt(0);pf.widow_control=True
for st in ['Title','Heading 1','Heading 2']:
    s=doc.styles[st];font(s,'SimHei',18 if st=='Title' else (14 if st=='Heading 1' else 12))
    s.paragraph_format.first_line_indent=Pt(0)
    s.paragraph_format.line_spacing=Pt(22 if st!='Title' else 28)
    s.paragraph_format.keep_with_next=True;s.paragraph_format.keep_together=True
    s.paragraph_format.space_before=Pt(10 if st!='Title' else 0)
    s.paragraph_format.space_after=Pt(6 if st!='Title' else 12)
font(doc.styles['Footer'],'SimSun',10)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='1';rr.append(tt);field.append(rr);footer._p.append(field)

def table(rows):
    coln=len(rows[0]);tb=doc.add_table(rows=1, cols=coln);tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    widths=[2.1,6.95,6.95] if coln==3 else [3.0,13.0]
    for column,width in zip(tb.columns,widths): column.width=Cm(width)
    pr=tb._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for tag,value in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
        e=OxmlElement('w:'+tag);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for i,row in enumerate(rows):
        cells=tb.rows[0].cells if i==0 else tb.add_row().cells
        trPr=cells[0]._tc.getparent().get_or_add_trPr()
        cannot=OxmlElement('w:cantSplit');trPr.append(cannot)
        if i==0: repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for j,text in enumerate(row):
            cell=cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=Pt(22);p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==0 or coln==3 else WD_ALIGN_PARAGRAPH.LEFT
            run=p.add_run(text);run.bold=i==0
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7');cell._tc.get_or_add_tcPr().append(shade)
    after=doc.add_paragraph();after.paragraph_format.space_after=Pt(0);after.paragraph_format.line_spacing=Pt(4)

def paragraph_with_links(value):
    p=doc.add_paragraph()
    pattern=re.compile(r'\[([^\]]+)\]\((https?://[^)]+)\)')
    last=0
    for match in pattern.finditer(value):
        p.add_run(value[last:match.start()])
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(match.group(2),RT.HYPERLINK,is_external=True))
        rr=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'2563A0');rp.append(co);rr.append(rp);tt=OxmlElement('w:t');tt.text=match.group(1);rr.append(tt);link.append(rr);p._p.append(link);last=match.end()
    p.add_run(value[last:])
    return p

lines=source.read_text().splitlines(); i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',x) for x in vals):rows.append(vals)
            i+=1
        table(rows);continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif re.match(r'^[一二三四五六七八九十]+、',line):doc.add_paragraph(line,'Heading 1')
    else:paragraph_with_links(line)
    i+=1

doc.core_properties.title='专利情报学习原型产品与技术说明'
doc.core_properties.subject='产品范围与共用核心实现机制'
doc.core_properties.author='专利情报AI项目'
doc.core_properties.keywords='专利情报,产品说明,数据契约,技术架构'
doc.save(DOCX)
print(DOCX)
